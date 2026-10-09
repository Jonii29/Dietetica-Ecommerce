"""Lectura "inteligente" de listas de precios: Excel, CSV o PDF con cualquier formato.

La idea es que la dueña pueda subir la lista tal como la tiene (la del proveedor, la que arma en
Excel, una exportación del sistema de caja) sin adaptarla a la plantilla:

* Busca la fila de encabezados en las primeras filas (puede haber logo, título, fecha arriba).
* Reconoce columnas por nombre aunque se llamen distinto: "Descripción", "Artículo", "P. Venta",
  "Precio público", "Rubro", "Cód.", "Peso"...
* Si hay varias columnas de precio elige la de venta al público (ignora costo, mayorista, sin IVA).
* Filas de sección ("FRUTOS SECOS" sin precio) pasan a ser la categoría de las filas siguientes.
  En un Excel con una hoja por rubro, el nombre de la hoja es la categoría.
* Del nombre saca la presentación ("Almendras x 500 g" -> producto "Almendras", 500 g) y las
  etiquetas dietarias ("sin TACC", "vegano", "orgánico", "keto", "sin azúcar").
* Si no hay código, genera uno estable a partir del nombre y la presentación, así una nueva
  lista con precios actualizados pisa los mismos productos en vez de duplicarlos.

El resultado son filas con las columnas canónicas de la plantilla, listas para `import_catalog`,
más un resumen de cómo se interpretó el archivo para mostrarlo en el panel.
"""
from __future__ import annotations

import csv
import hashlib
import io
import re
from dataclasses import dataclass, field

from ..utils.text import normalize, slugify

# ---------------------------------------------------------------------------
# Campos y sinónimos de encabezados
# ---------------------------------------------------------------------------

# Etiquetas visibles en el panel para cada campo
FIELD_LABELS: dict[str, str] = {
    "sku": "Código / SKU",
    "producto": "Agrupar como producto",
    "nombre": "Nombre del producto",
    "marca": "Marca",
    "categoria": "Categoría",
    "subcategoria": "Subcategoría",
    "variante": "Presentación (texto)",
    "tipo_venta": "Tipo de venta",
    "contenido": "Contenido / peso",
    "precio": "Precio de venta",
    "precio_lista": "Precio anterior (tachado)",
    "stock": "Stock",
    "etiquetas": "Etiquetas",
    "sinonimos": "Sinónimos",
    "descripcion": "Descripción",
    "imagen_url": "Imagen (URL)",
    "activo": "Activo (sí/no)",
    "orden": "Orden",
}

# Frases (ya normalizadas: minúsculas, sin tildes) que identifican cada campo.
# Coincidencia exacta suma más que coincidencia por palabra contenida.
_ALIASES: dict[str, list[str]] = {
    "sku": ["sku", "codigo", "cod", "cod art", "codigo articulo", "cod producto", "id", "art",
            "nro", "numero", "ref", "referencia", "plu", "codigo interno"],
    "producto": ["producto padre", "grupo producto", "agrupar"],
    "nombre": ["nombre", "producto", "descripcion", "articulo", "detalle", "item", "denominacion",
               "nombre producto", "descripcion producto", "mercaderia"],
    "marca": ["marca", "fabricante", "laboratorio", "elaborador"],
    "categoria": ["categoria", "rubro", "familia", "seccion", "linea", "grupo", "tipo producto",
                  "departamento"],
    "subcategoria": ["subcategoria", "subrubro", "sub rubro", "subfamilia", "sub familia",
                     "sub categoria"],
    "variante": ["variante", "presentacion", "formato", "envase"],
    "tipo_venta": ["tipo venta", "venta por", "modalidad"],
    "contenido": ["contenido", "peso", "gramaje", "tamano", "medida", "gramos", "kg", "peso neto",
                  "contenido neto", "cantidad contenido", "volumen"],
    "precio": ["precio", "precio venta", "pvp", "precio publico", "precio final", "venta",
               "importe", "valor", "precio unitario", "p venta", "pv", "publico", "minorista",
               "precio minorista", "precio con iva", "precio lista", "lista"],
    "precio_lista": ["precio anterior", "antes", "precio tachado", "precio sin descuento",
                     "precio regular", "precio original"],
    "stock": ["stock", "existencia", "existencias", "disponible", "cantidad", "cant", "unidades",
              "stock actual"],
    "etiquetas": ["etiquetas", "tags", "atributos", "caracteristicas", "apto"],
    "sinonimos": ["sinonimos", "palabras clave", "keywords", "busqueda"],
    "descripcion": ["descripcion larga", "detalle largo", "observaciones", "comentario", "notas",
                    "info"],
    "imagen_url": ["imagen", "imagen url", "foto", "url imagen", "image", "imagen_url"],
    "activo": ["activo", "habilitado", "publicado", "visible", "mostrar"],
    "orden": ["orden", "posicion"],
}

# Columnas de precio que NO son el precio de venta al público
_PRICE_EXCLUDE = re.compile(r"\b(costo|compra|mayorista|mayor|sin iva|neto|proveedor|margen|"
                            r"ganancia|markup|dolar|usd|u s s|bulto|caja)\b")
_PRICE_PREFER = re.compile(r"\b(venta|publico|final|minorista|pvp|con iva)\b")

# Columnas con una etiqueta dietaria como encabezado y "x"/"sí" como valor
_TAG_COLUMNS = {
    "sin tacc": "sin tacc", "sin gluten": "sin tacc", "apto celiaco": "sin tacc", "celiaco": "sin tacc",
    "vegano": "vegano", "vegan": "vegano", "organico": "organico", "keto": "keto",
    "sin azucar": "sin azucar", "sin lactosa": "sin lactosa",
}
_TAG_IN_TEXT = [
    (re.compile(r"\b(sin t ?a ?c ?c|sin gluten|gluten free|apto celiac)"), "sin tacc"),
    (re.compile(r"\b(vegan[oa]?s?|plant based)\b"), "vegano"),
    (re.compile(r"\b(organic[oa]s?|organic)\b"), "organico"),
    (re.compile(r"\bketo\b"), "keto"),
    (re.compile(r"\b(sin azucar|0 azucar|zero azucar|sugar free)"), "sin azucar"),
    (re.compile(r"\b(sin lactosa|deslactosad)"), "sin lactosa"),
]

_TRUTHY = {"x", "si", "s", "1", "true", "verdadero", "yes", "y", "ok", "✓", "✔"}

# Presentación dentro del nombre: "x 500 g", "500grs", "1 kg", "1,5 L", "x 6 u", "360 cc"
_SIZE_RE = re.compile(
    r"(?:\bx\s*)?(\d+(?:[.,]\d+)?)\s*(kgs?|kilos?|grs?|gramos|g|ml|cc|lts?|litros?|l|u|un|unid|unidades)\b\.?",
    re.IGNORECASE,
)
_BULK_RE = re.compile(r"\b(granel|suelto|fraccionado|por kilo|x kilo|x kg|el kilo)\b")


@dataclass
class ColumnInfo:
    index: int
    header: str
    field: str | None
    examples: list[str]


@dataclass
class ParsedSheet:
    rows: list[dict[str, str]]                       # filas canónicas (+ "_line")
    columns: list[ColumnInfo] = field(default_factory=list)
    sheets: list[str] = field(default_factory=list)  # hojas usadas
    header_row: int | None = None                    # 1-based en el archivo; None = sin encabezado
    notes: list[str] = field(default_factory=list)   # avisos legibles para el panel
    skipped: int = 0                                 # filas vacías / títulos ignorados

    def summary(self) -> dict:
        return {
            "sheets": self.sheets,
            "header_row": self.header_row,
            "notes": self.notes,
            "skipped": self.skipped,
            "columns": [
                {"index": c.index, "header": c.header, "field": c.field, "examples": c.examples}
                for c in self.columns
            ],
            "fields": FIELD_LABELS,
        }


# ---------------------------------------------------------------------------
# 1) Archivo -> tablas de celdas
# ---------------------------------------------------------------------------

def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def read_tables(data: bytes, filename: str) -> list[tuple[str, list[list[str]]]]:
    """Devuelve [(nombre_hoja, filas)] para xlsx/xlsm, xls, csv/tsv/txt o pdf."""
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        out = []
        for ws in wb.worksheets:
            if getattr(ws, "sheet_state", "visible") != "visible":
                continue
            out.append((ws.title, [[_cell(v) for v in r] for r in ws.iter_rows(values_only=True)]))
        wb.close()
        return out
    if name.endswith(".xls"):
        import xlrd  # Excel 97-2003

        book = xlrd.open_workbook(file_contents=data)
        return [(sh.name, [[_cell(v) for v in sh.row_values(i)] for i in range(sh.nrows)])
                for sh in book.sheets() if sh.visibility == 0]
    if name.endswith(".pdf"):
        return _read_pdf(data)

    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("latin-1")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = [[c.strip() for c in r] for r in csv.reader(io.StringIO(text), dialect=dialect)]
    return [("", rows)]


def _read_pdf(data: bytes) -> list[tuple[str, list[list[str]]]]:
    """Tablas de un PDF (listas de proveedores). Si no hay tablas dibujadas, separa el texto
    de cada línea por espacios largos."""
    import pdfplumber

    rows: list[list[str]] = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            tables = page.extract_tables()
            if tables:
                for t in tables:
                    rows.extend([[_cell(c).replace("\n", " ") for c in r] for r in t])
            else:
                for line in (page.extract_text() or "").splitlines():
                    parts = [p.strip() for p in re.split(r"\s{2,}|\t", line) if p.strip()]
                    if len(parts) == 1:
                        # "Almendras x 500g $ 8.100" -> nombre + precio al final
                        m = re.match(r"^(.*?)[\s.]*\$?\s*(\d{1,3}(?:[.,]\d{3})*(?:[.,]\d{1,2})?)$", line.strip())
                        parts = [m.group(1), m.group(2)] if m else parts
                    rows.append(parts)
    return [("", rows)]


# ---------------------------------------------------------------------------
# 2) Detección de encabezado y columnas
# ---------------------------------------------------------------------------

def _score_header(h: str, alias_list: list[str]) -> float:
    if not h:
        return 0.0
    best = 0.0
    words = set(h.split())
    for a in alias_list:
        if h == a:
            return 3.0
        a_words = a.split()
        if all(w in words for w in a_words):
            best = max(best, 1.5 + 0.2 * len(a_words))
        elif len(a) >= 4 and h.startswith(a):
            best = max(best, 1.0)
    return best


def map_columns(headers: list[str]) -> dict[int, str]:
    """Asigna a cada columna su campo más probable sin repetir campos."""
    norm = [normalize(h) for h in headers]
    candidates: list[tuple[float, int, str]] = []
    for i, h in enumerate(norm):
        if not h:
            continue
        for f, aliases in _ALIASES.items():
            # El nombre exacto de la plantilla siempre gana ("producto", "precio_lista"...)
            s = 4.0 if h == f.replace("_", " ") else _score_header(h, aliases)
            if f == "precio" and s:
                if _PRICE_EXCLUDE.search(h):
                    s = 0.3  # solo si no hay otra
                elif _PRICE_PREFER.search(h):
                    s += 1.0
            if f == "precio_lista" and s and _PRICE_EXCLUDE.search(h):
                s = 0
            if f in ("contenido", "stock") and "precio" in h.split():
                s = 0  # "Precio x kg" es un precio, no un peso
            if s:
                candidates.append((s, i, f))
    # Mayor puntaje primero; a igualdad, la columna más a la izquierda
    candidates.sort(key=lambda t: (-t[0], t[1]))
    taken_cols: set[int] = set()
    result: dict[int, str] = {}
    for s, i, f in candidates:
        if i in taken_cols or f in result.values():
            continue
        result[i] = f
        taken_cols.add(i)
    # Listas comunes: "Producto" o "Descripción" es el nombre; "Precio lista" es el de venta.
    fields = set(result.values())
    for alt, target in (("producto", "nombre"), ("descripcion", "nombre"), ("precio_lista", "precio")):
        if target not in fields and alt in fields:
            col = next(i for i, f in result.items() if f == alt)
            result[col] = target
            fields = set(result.values())
    return result


def _find_header(rows: list[list[str]]) -> tuple[int | None, dict[int, str]]:
    best: tuple[float, int | None, dict[int, str]] = (0.0, None, {})
    for idx, row in enumerate(rows[:30]):
        filled = [c for c in row if c]
        if len(filled) < 2:
            continue
        # Un encabezado tiene texto, no precios
        if sum(_looks_like_price(c) for c in filled) > len(filled) / 2:
            continue
        mapping = map_columns(row)
        fields = set(mapping.values())
        score = len(fields) + (2 if "nombre" in fields else 0) + (2 if "precio" in fields else 0)
        if score > best[0]:
            best = (score, idx, mapping)
    if best[1] is None or not {"nombre", "precio"} <= set(best[2].values()):
        return None, {}
    return best[1], best[2]


def _looks_like_price(s: str) -> bool:
    s = s.strip()
    return bool(re.fullmatch(r"\$?\s*-?\d{1,3}(?:[.,]?\d{3})*(?:[.,]\d{1,2})?\s*", s)) and any(ch.isdigit() for ch in s)


def _guess_without_header(rows: list[list[str]]) -> dict[int, str]:
    """Sin encabezado: la columna con más texto es el nombre y la de más números con pinta
    de precio (la de más a la derecha si hay varias) es el precio."""
    width = max((len(r) for r in rows), default=0)
    text_score = [0] * width
    price_score = [0] * width
    for r in rows[:200]:
        for i, c in enumerate(r):
            if not c:
                continue
            if _looks_like_price(c):
                price_score[i] += 1
            elif re.search(r"[a-zA-Z]{3,}", c):
                text_score[i] += len(c)
    if not width or max(text_score) == 0 or max(price_score) == 0:
        return {}
    name_col = text_score.index(max(text_score))
    top = max(price_score)
    price_col = max(i for i, s in enumerate(price_score) if s >= top * 0.8 and i != name_col)
    return {name_col: "nombre", price_col: "precio"}


# ---------------------------------------------------------------------------
# 3) Filas -> filas canónicas
# ---------------------------------------------------------------------------

def split_size(name: str) -> tuple[str, str | None]:
    """'Almendras peladas x 500 grs.' -> ('Almendras peladas', '500g')."""
    matches = list(_SIZE_RE.finditer(name))
    if not matches:
        return name.strip(), None
    m = matches[-1]
    qty, unit = m.group(1), m.group(2).lower()
    unit = {"kgs": "kg", "kilo": "kg", "kilos": "kg", "grs": "g", "gr": "g", "gramos": "g",
            "lts": "l", "lt": "l", "litro": "l", "litros": "l", "un": "u", "unid": "u",
            "unidades": "u"}.get(unit, unit)
    base = (name[: m.start()] + name[m.end():]).strip(" -–,.(")
    base = re.sub(r"\s+x$", "", re.sub(r"\s{2,}", " ", base)).strip()
    return (base or name.strip()), f"{qty}{unit}"


def stable_sku(product_key: str, label: str) -> str:
    """Código estable para filas sin código: mismo nombre + presentación = mismo SKU."""
    digest = hashlib.sha1(f"{product_key}|{slugify(label)}".encode()).hexdigest()[:8].upper()
    prefix = re.sub(r"[^A-Z]", "", slugify(product_key).upper())[:3] or "PRD"
    return f"{prefix}-{digest}"


def detect_tags(*texts: str) -> list[str]:
    blob = normalize(" ".join(t for t in texts if t))
    return [tag for rx, tag in _TAG_IN_TEXT if rx.search(blob)]


def parse_tables(
    tables: list[tuple[str, list[list[str]]]],
    override: dict[int, str] | None = None,
) -> ParsedSheet:
    out = ParsedSheet(rows=[])
    multi = len([t for t in tables if any(any(c for c in r) for r in t[1])]) > 1
    line_offset = 0

    for sheet_name, rows in tables:
        if not any(any(c for c in r) for r in rows):
            continue
        if override:
            header_idx = _find_header(rows)[0]
            mapping = {i: f for i, f in override.items() if f}
            if header_idx is None:
                # Si la primera fila con datos no tiene un precio en la columna elegida, es el encabezado
                pcol = next((i for i, f in mapping.items() if f == "precio"), None)
                k = next((j for j, r in enumerate(rows) if any(r)), None)
                if k is not None and pcol is not None and not _looks_like_price(rows[k][pcol] if pcol < len(rows[k]) else ""):
                    header_idx = k
        else:
            header_idx, mapping = _find_header(rows)
            if not mapping:
                mapping = _guess_without_header(rows)
                header_idx = None
                if mapping:
                    out.notes.append("No encontré encabezados: tomé la columna con texto como nombre "
                                     "y la de números como precio. Revisá las columnas abajo.")
        if not mapping or "nombre" not in mapping.values() or "precio" not in mapping.values():
            if multi:
                out.notes.append(f"Hoja «{sheet_name}» ignorada: no encontré columnas de nombre y precio.")
                continue
            # Mostrar las columnas para que la dueña elija cuál es cuál
            k = next((j for j, r in enumerate(rows) if any(r)), 0)
            width = max(len(r) for r in rows)
            for i in range(width):
                ex = [r[i] for r in rows[k + 1:] if i < len(r) and r[i]][:3]
                h = rows[k][i] if i < len(rows[k]) else ""
                if h or ex:
                    out.columns.append(ColumnInfo(i, h or f"Columna {i + 1}", mapping.get(i), ex))
            out.header_row = k + 1
            out.sheets.append(sheet_name or "Hoja 1")
            out.notes.append("No pude reconocer qué columna es el nombre y cuál el precio. "
                             "Elegilas en la tabla de columnas y volvé a revisar.")
            break

        if not out.columns:  # se muestran las columnas de la primera hoja usada
            headers = rows[header_idx] if header_idx is not None else []
            width = max(len(r) for r in rows)
            data_rows = rows[(header_idx + 1) if header_idx is not None else 0:]
            for i in range(width):
                ex = [r[i] for r in data_rows if i < len(r) and r[i] and sum(1 for c in r if c) > 1][:3]
                h = headers[i] if i < len(headers) else ""
                if h or ex:
                    out.columns.append(ColumnInfo(i, h or f"Columna {i + 1}", mapping.get(i), ex))
            out.header_row = None if header_idx is None else line_offset + header_idx + 1

        headers_norm = {i: normalize(rows[header_idx][i]) for i in range(len(rows[header_idx]))} \
            if header_idx is not None else {}
        price_header = next((headers_norm.get(i, "") for i, f in mapping.items() if f == "precio"), "")
        price_per_kg = bool(re.search(r"\b(kg|kilo)\b", price_header))
        content_header = next((headers_norm.get(i, "") for i, f in mapping.items() if f == "contenido"), "")
        content_unit = "kg" if re.search(r"\b(kg|kilo)", content_header) else (
            "ml" if re.search(r"\b(ml|cc)\b", content_header) else "g")
        tag_cols = {i: tag for i, h in headers_norm.items() if i not in mapping and h in _TAG_COLUMNS
                    for tag in [_TAG_COLUMNS[h]]}

        out.sheets.append(sheet_name or "Hoja 1")
        section: str | None = sheet_name if multi and sheet_name else None
        start = (header_idx + 1) if header_idx is not None else 0

        for ridx in range(start, len(rows)):
            r = rows[ridx]
            line = line_offset + ridx + 1
            get = lambda f: next((r[i] for i, ff in mapping.items() if ff == f and i < len(r)), "")  # noqa: E731
            name, price = get("nombre").strip(), get("precio").strip()
            filled = [c for c in r if c]
            if not filled:
                continue
            if not price and len(filled) == 1 and not _looks_like_price(filled[0]) \
                    and re.search(r"[a-zA-Z]{3,}", filled[0]):
                # Fila de sección: "FRUTOS SECOS" (en cualquier columna)
                title = filled[0].strip(" :-")
                section = title.title() if title.isupper() else title
                out.skipped += 1
                continue
            if headers_norm and normalize(name) == headers_norm.get(
                    next(i for i, f in mapping.items() if f == "nombre"), None):
                out.skipped += 1  # encabezado repetido (ej. en cada página del PDF)
                continue
            if not name and not price:
                out.skipped += 1
                continue

            row = {f: (r[i].strip() if i < len(r) else "") for i, f in mapping.items()}
            row["_line"] = str(line)

            # Presentación
            if row.get("contenido") and not re.fullmatch(r"\d+(?:[.,]\d+)?", row["contenido"]):
                row["contenido"] = split_size(row["contenido"])[1] or row["contenido"]
            if not row.get("contenido") and row.get("variante"):
                row["contenido"] = split_size(row["variante"])[1] or ""
            if not row.get("contenido"):
                base, size = split_size(name)
                if size:
                    row["contenido"] = size
                    if not row.get("producto"):
                        row["nombre"] = base
                elif price_per_kg or _BULK_RE.search(normalize(name)):
                    row["contenido"] = "1kg"
                    row.setdefault("tipo_venta", "peso")
            elif re.fullmatch(r"\d+(?:[.,]\d+)?", row["contenido"]):
                row["contenido"] += content_unit
            if price_per_kg or _BULK_RE.search(normalize(name)):
                row["tipo_venta"] = row.get("tipo_venta") or "peso"
                row["nombre"] = re.sub(r"\s*\b(a granel|granel|suelto|x kg|por kilo)\b", "",
                                       row["nombre"], flags=re.I).strip() or row["nombre"]

            # Categoría por sección / hoja / sin dato
            if not row.get("categoria"):
                row["categoria"] = section or "Varios"
                if not section:
                    row["_categoria_por_defecto"] = "1"

            # Etiquetas: columna propia + columnas tipo "Sin TACC: x" + texto del nombre
            if "etiquetas" in row or tag_cols:
                row["_etiquetas_completas"] = "1"
            tags = [t for t in re.split(r"[|;,/]", row.get("etiquetas", "")) if t.strip()]
            tags += [tag for i, tag in tag_cols.items() if i < len(r) and normalize(r[i]) in _TRUTHY]
            tags += detect_tags(name, row.get("descripcion", ""))
            row["etiquetas"] = "|".join(dict.fromkeys(t.strip() for t in tags))

            # Stock con texto
            st = normalize(row.get("stock", ""))
            if "stock" not in row:
                pass
            elif st in ("", "-", "consultar", "s d", "sd"):
                row["stock"] = ""
            elif st in ("sin stock", "agotado", "no", "no hay"):
                row["stock"] = "0"
            elif not re.fullmatch(r"-?\d+(?:[.,]\d+)?", row.get("stock", "").replace(".", "")):
                row["stock"] = ""

            out.rows.append(row)
        line_offset += len(rows)

    if override:
        cols = {c.index: c for c in out.columns}
        for i, f in override.items():
            if i in cols:
                cols[i].field = f or None
    return out


def finalize_skus(rows: list[dict[str, str]]) -> None:
    """Completa SKUs faltantes (estables) y desambigua duplicados generados."""
    used: set[str] = {r["sku"].strip().upper() for r in rows if r.get("sku")}
    for r in rows:
        if r.get("sku"):
            continue
        key = slugify(r.get("producto") or f"{r.get('nombre', '')} {r.get('marca', '')}")
        label = r.get("variante") or r.get("contenido") or ""
        sku, n = stable_sku(key, label), 2
        while sku in used:
            sku, n = f"{stable_sku(key, label)}-{n}", n + 1
        used.add(sku)
        r["sku"] = sku
        r["_sku_auto"] = "1"
